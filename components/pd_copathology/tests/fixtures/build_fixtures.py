"""Build a completely synthetic stand-in for a PPMI download directory.

Nothing here is derived from PPMI data. Participant identifiers, dates, values
and free text are invented, and deliberately recognisable so tests can assert
that none of them reach a report:

* participant identifiers are 9900001-9900080 (61-80 exist only in remote-only tables);
* birth dates are in 1931, visit dates in 2031, diagnosis and medication dates in 2030/2032;
* every free-text cell starts with ``SYNTHTEXT``.

Usage::

    python tests/fixtures/build_fixtures.py [destination]
"""

from __future__ import annotations

import csv
import io
import os
import shutil
import sys
import zipfile
from pathlib import Path

DEFAULT_DEST = Path(__file__).resolve().parent / "synthetic_ppmi"

PARTICIPANT_IDS = [str(9900000 + i) for i in range(1, 81)]
CLINIC_IDS = PARTICIPANT_IDS[:60]  # the ones present in Participant_Status
FREE_TEXT_MARKER = "SYNTHTEXT"
DATE_MARKERS = ["/1931", "/2031", "2031-", "/2030", "/2032"]
_STAMP_A = (2024, 1, 1, 9, 0, 0)
_STAMP_B = (2024, 3, 15, 9, 0, 0)


def _csv(header: list[str], rows: list[list[object]]) -> bytes:
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(header)
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8")


def _zip(entries: list[tuple[str, bytes]], stamp: tuple) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in entries:
            zf.writestr(zipfile.ZipInfo(name, date_time=stamp), data, compress_type=zipfile.ZIP_DEFLATED)
    return buffer.getvalue()


def _pid(i: int) -> str:
    return PARTICIPANT_IDS[i - 1]


def _cohort(i: int) -> tuple[str, str]:
    if i <= 25 or i > 50:
        return "1", "Parkinson's Disease"
    if i <= 40:
        return "2", "Healthy Control"
    if i <= 47:
        return "4", "Prodromal"
    return "3", "SWEDD"


def participant_status(n: int) -> bytes:
    rows = []
    for i in range(1, n + 1):
        code, label = _cohort(i)
        rows.append([_pid(i), code, label, f"{(i % 12) + 1:02d}/2031", "Enrolled", 1 if i % 7 == 0 else 0])
    return _csv(["PATNO", "COHORT", "COHORT_DEFINITION", "ENROLL_DATE", "ENROLL_STATUS", "ENRLLRRK2"], rows)


def demographics() -> bytes:
    rows = [[1000 + i, _pid(i), "SC", i % 2, f"{(i % 12) + 1:02d}/1931", "" if i % 9 == 0 else (i % 3) + 1]
            for i in range(1, 51)]
    return _csv(["REC_ID", "PATNO", "EVENT_ID", "SEX", "BIRTHDT", "HANDED"], rows)


def updrs3(extra_column: bool) -> bytes:
    header = ["REC_ID", "PATNO", "EVENT_ID", "PAG_NAME", "INFODT", "PDSTATE", "NP3RIGN", "NP3FTAPR", "NP3TOT", "NHY"]
    if extra_column:
        header.insert(6, "NP3SPCH")
    rows, rec = [], 2000

    def add(i: int, event: str, state: str) -> None:
        nonlocal rec
        rec += 1
        total = "" if (i + len(event)) % 11 == 0 else 10 + (i * 3) % 40
        nhy = "UR" if i % 13 == 0 else (i % 5)
        row = [rec, _pid(i), event, "NUPDRS3", "02/2031", state, i % 5, (i + 1) % 5, total, nhy]
        if extra_column:
            row.insert(6, i % 4)
        rows.append(row)

    for i in range(1, 41):
        add(i, "BL", "")
        add(i, "V04", "ON" if i <= 10 else "")
        if i <= 10:
            add(i, "V04", "OFF")  # second record for the same participant and visit
        if i <= 20:
            add(i, "V08", "")
    return _csv(header, rows)


def moca() -> bytes:
    rows, rec = [], 3000
    for i in range(1, 46):
        for event in ("SC", "V04"):
            rec += 1
            rows.append([rec, _pid(i), event, "03/2031", "" if i % 10 == 0 else 18 + i % 12])
    return _csv(["REC_ID", "PATNO", "EVENT_ID", "INFODT", "MCATOT"], rows)


_BIO_HEADER = ["PATNO", "SEX", "COHORT", "CLINICAL_EVENT", "TYPE", "TESTNAME", "TESTVALUE", "UNITS",
               "RUNDATE", "PROJECTID", "PI_NAME"]


def _bio_row(i: int, event: str, matrix: str, test: str, value: object, units: str, project: int) -> list:
    return [_pid(i), i % 2, _cohort(i)[0], event, matrix, test, value, units, "2031-04-02", project, "SYNTHTEXT Lab"]


def current_biospecimen() -> bytes:
    rows = []
    for i in range(1, 31):
        # Abeta42: 6 x "<LLOQ", 2 x ">ULOQ", 5 x "ND", 5 x "-1", 1 blank, the rest numeric
        if i <= 6:
            abeta: object = "<LLOQ"
        elif i <= 8:
            abeta = ">ULOQ"
        elif i <= 13:
            abeta = "ND"
        elif i <= 18:
            abeta = -1
        elif i == 19:
            abeta = ""
        else:
            abeta = 600 + i * 7.5
        csf = "Cerebrospinal Fluid"
        rows.append(_bio_row(i, "BL", csf, "ABeta 1-42", abeta, "pg/mL", 125))
        rows.append(_bio_row(i, "BL", csf, "pTau", 12 + i * 0.4, "pg/mL", 125))
        rows.append(_bio_row(i, "BL", csf, "tTau", 150 + i * 2, "pg/mL", 125))
        if i <= 15:
            rows.append(_bio_row(i, "V04", csf, "ABeta 1-42", 610 + i * 7, "pg/mL", 125))
            rows.append(_bio_row(i, "V04", csf, "pTau", 13 + i * 0.4, "pg/mL", 125))
        if i <= 20:
            rows.append(_bio_row(i, "BL", csf, "CSF Alpha-synuclein", 1500 + i * 11, "pg/mL", 124))
        if i <= 12:
            rows.append(_bio_row(i, "BL", "Plasma", "pTau217", 0.2 + i * 0.01, "pg/mL", 300))
        if i <= 25:
            rows.append(_bio_row(i, "BL", "Serum", "NfL", "<LLOQ" if i <= 5 else 9 + i * 0.3, "pg/mL", 250))
    for i in range(1, 23):
        matrix = "Cerebrospinal Fluid" if i <= 20 else "CSF"
        rows.append(_bio_row(i, "BL", matrix, "ABeta42", 700 + i * 6, "pg/mL", 283))
        if i <= 18:
            rows.append(_bio_row(i, "BL", matrix, "pTau181", 15 + i * 0.3, "pg/mL", 283))
    for i in range(1, 36):
        rows.append(_bio_row(i, "SC", "DNA", "APOE GENOTYPE", ["E3/E3", "E3/E4", "E4/E4", "E2/E3"][i % 4], "", 118))
    for i in range(1, 4):
        rows.append(_bio_row(i, "BL", "Plasma", "Rare Analyte X", 3.3, "AU", 301))
    rows.append(_bio_row(1, "BL", "Cerebrospinal Fluid", "pTau", 12.9, "pg/mL", 125))  # repeated measurement
    for i in range(10, 31):
        for test, value in (("ABeta", 650 + i), ("ABeta raw", 640 + i), ("pTau", 14 + i * 0.2), ("tTau", 160 + i)):
            rows.append(_bio_row(i, "BL", "Cerebrospinal Fluid", test, value, "pg/mL", 159))
    for i in range(1, 9):
        rows.append(_bio_row(i, "BL", "Plasma", "p217+tau", 0.1 + i * 0.01, "pg/ml", 208))
    for i in range(5, 17):
        rows.append(_bio_row(i, "V04", "PLASMA" if i == 5 else "Plasma", "PTAU217", 0.3 + i * 0.01, "pg/mL", 292))
    return _csv(_BIO_HEADER, rows)


def pilot_biospecimen() -> bytes:
    rows = [_bio_row(i, "BL", "Cerebrospinal Fluid", "ABeta 1-42", 0.6 + i * 0.01, "ng/L", 999) for i in range(1, 11)]
    return _csv(_BIO_HEADER, rows)


def code_list() -> bytes:
    rows = [["PATIENT_STATUS", "COHORT", "COHORT", c, d] for c, d in
            [("1", "Parkinson's Disease"), ("2", "Healthy Control"), ("3", "SWEDD"), ("4", "Prodromal")]]
    rows += [["ALL", "EVENT_ID", "EVENT_ID", c, d] for c, d in
             [("SC", "Screening"), ("BL", "Baseline"), ("V04", "Month 12"), ("V08", "Month 24"), ("LOG", "Log form"),
              ("ED", "Event Driven")]]
    rows += [["MHXGEN", "MHCAT", "MHCAT", c, d] for c, d in
             [("106", "Cardiovascular"), ("112", "Metabolic/Endocrine"), ("116", "Allergy/Immunologic")]]
    return _csv(["MOD_NAME", "ITM_NAME", "CDL_NAME", "CODE", "DECODE"], rows)


def data_dictionary() -> bytes:
    rows = [["NUPDRS3", "NP3TOT", "MDS-UPDRS Part III total score", "NUMBER"],
            ["MOCA", "MCATOT", "MoCA total score", "NUMBER"],
            ["VITAL", "SYSSUP", "Supine systolic blood pressure", "NUMBER"]]
    return _csv(["MOD_NAME", "ITM_NAME", "DSCR", "ITM_TYPE"], rows)


def genetics() -> bytes:
    rows = [[_pid(i), ["E3/E3", "E3/E4", "E4/E4", "E2/E3"][i % 4], i % 9 == 0, i % 7 == 0, 0, i % 3]
            for i in range(1, 36)]
    return _csv(["PATNO", "APOE", "GBA", "LRRK2", "SNCA", "PATHVAR_COUNT"], rows)


def vital_signs() -> bytes:
    rows, rec = [], 4000
    for i in range(1, 41):
        for event in ("BL", "V04"):
            rec += 1
            rows.append([rec, _pid(i), event, "05/2031", 60 + i % 30, 160 + i % 25, 110 + i % 40,
                         "" if i % 8 == 0 else 65 + i % 25])
    for i in range(41, 45):  # no baseline measurement: the screening one is used instead
        rec += 1
        rows.append([rec, _pid(i), "SC", "04/2031", 70, 170, 118, 76])
    return _csv(["REC_ID", "PATNO", "EVENT_ID", "INFODT", "WGTKG", "HTCM", "SYSSUP", "DIASUP"], rows)


def medical_conditions() -> bytes:
    terms = ["essential hypertension", "type 2 diabetes", "seasonal allergy", "hyperlipidemia"]
    categories = [106, 112, 116, 112]
    rows, rec = [], 5000
    for i in range(1, 31):
        year = 2030 if i % 2 == 0 else 2032  # baseline visits are in 2031
        for k in range(1 + i % 2):
            rec += 1
            rows.append([rec, _pid(i), "ED", categories[(i + k) % 4],
                         f"{FREE_TEXT_MARKER} {terms[(i + k) % 4]} noted by clinician", f"06/{year}", year])
    return _csv(["REC_ID", "PATNO", "EVENT_ID", "MHCAT", "MHTERM", "MHDIAGDT", "MHDIAGYR"], rows)


def medications() -> bytes:
    drugs = ["atorvastatin 20mg", "lisinopril 10mg", "vitamin d", "metformin 500mg"]
    rows, rec = [], 6000
    reasons = ["high cholesterol", "hypertension", "supplement", "diabetes"]
    for i in range(1, 31):
        rec += 1
        year = 2030 if i % 3 == 0 else 2032
        stop = "08/2030" if i % 6 == 0 else ""  # started and stopped before the 2031 baseline
        rows.append([rec, _pid(i), "LOG", f"{FREE_TEXT_MARKER} {drugs[i % 4]}", f"06/{year}", stop, 0 if stop else 1,
                     1 + i % 3, f"{FREE_TEXT_MARKER} for {reasons[i % 4]}", drugs[i % 4].split()[0].upper()])

    def extra(i: int, drug: str, reason: str) -> None:
        nonlocal rec
        rec += 1
        rows.append([rec, _pid(i), "LOG", f"{FREE_TEXT_MARKER} {drug} tablets", "06/2030", "", 1, 2,
                     f"{FREE_TEXT_MARKER} {reason}" if reason else "", ""])

    for i in range(31, 37):
        extra(i, "Propranolol", "essential tremor")
    for i in (1, 2):
        extra(i, "propranolol", "hypertension")
    for i in range(37, 42):
        extra(i, "Flomax", "enlarged prostate")  # brand name of tamsulosin
    for i in (42, 43, 44, 45):
        extra(i, "midodrine", "orthostatic hypotension")
    extra(3, "midodrine", "")
    for i in (4, 5):
        extra(i, "ramipril", "high blood pressure")  # fewer than five participants: collapsed
    for i in range(6, 13):
        extra(i, "amlodipine", "hypertension")
    return _csv(["REC_ID", "PATNO", "EVENT_ID", "CMTRT", "STARTDT", "STOPDT", "ONGOING", "CMINDC", "CMINDC_TEXT",
                 "WHODRUG"], rows)


def annotated_dictionary() -> bytes:
    rows = [["BIOANALYS", "TESTNAME", "Name of the biospecimen test", "CHAR", "synthetic mapping note"],
            ["BIOANALYS", "TESTVALUE", "Result of the biospecimen test", "CHAR", ""],
            ["CURATED", "fampd", "Family history of PD (curated)", "NUMBER", ""],
            ["CONMED", "CMTRT", "Medication", "CHAR", ""],
            ["CONMED", "STARTDT", "Start date", "DATE", ""],
            ["CONMED", "WHODRUG", "WHO drug name", "CHAR", ""],
            ["SOCIOECO", "EDUCYRS", "Years of education", "NUMBER", ""]]
    return _csv(["MOD_NAME", "ITM_NAME", "DSCR", "ITM_TYPE", "MAPPING_NOTES"], rows)


def fs7_aparc() -> bytes:
    rows = [[_pid(i), "BL", round(2.4 + i * 0.01, 2), round(2.5 + i * 0.01, 2)] for i in range(1, 21)]
    return _csv(["PATNO", "EVENT_ID", "lh_MeanThickness", "rh_MeanThickness"], rows)


def datscan() -> bytes:
    rows = [[_pid(i), "SC", "01/2031", round(1.5 + i * 0.03, 2), round(1.4 + i * 0.03, 2),
             "" if i % 12 == 0 else round(0.8 + i * 0.02, 2), round(0.7 + i * 0.02, 2)] for i in range(1, 39)]
    return _csv(["PATNO", "EVENT_ID", "DATSCAN_DATE", "DATSCAN_CAUDATE_R", "DATSCAN_CAUDATE_L",
                 "DATSCAN_PUTAMEN_R", "DATSCAN_PUTAMEN_L"], rows)


def blood_chemistry() -> bytes:
    rows = []
    for i in range(1, 21):
        rows.append([_pid(i), "SC", "Glucose", 80 + i, "mg/dL"])
        rows.append([_pid(i), "SC", "Cholesterol, Total", 160 + i * 2, "mg/dL"])
    return _csv(["PATNO", "EVENT_ID", "LTSTNAME", "LSIRES", "LSIUNIT"], rows)


def saa_workbook() -> bytes:
    from datetime import datetime

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SAA"
    sheet.append(["PATNO", "CLINICAL_EVENT", "SAA_Status", "SAAMethod"])
    for i in range(1, 23):
        sheet.append([int(_pid(i)), "BL", ["Positive", "Negative", "Positive", "Inconclusive"][i % 4], "24h"])
    readme = workbook.create_sheet("ReadMe")
    readme.append(["comment"])
    workbook.properties.created = datetime(2024, 3, 15)
    workbook.properties.modified = datetime(2024, 3, 15)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def curated_workbook() -> bytes:
    from datetime import datetime

    from openpyxl import Workbook

    workbook = Workbook()
    info = workbook.active
    info.title = "Information"
    info.append(["Revision Notes", "Date of Changes"])
    info.append(["Initial synthetic release of the curated cut", "2024-03-10"])
    info.append(["Added the fampd variable", "2024-03-12"])
    data = workbook.create_sheet("20240310")
    data.append(["PATNO", "COHORT", "EVENT_ID", "visit_date", "moca", "BMI", "fampd", "abeta", "ptau",
                 "IU_ABeta42_CSF", "MIA_PUTAMEN_L", "APOE", "ptau217_plasma"])
    for i in range(1, 46):
        csf = [600 + i, 14 + i * 0.1] if i <= 24 else [None, None]
        data.append([int(_pid(i)), int(_cohort(i)[0]), "BL", f"{(i % 12) + 1:02d}/2031", 20 + i % 10, 22 + i % 9, 1 + i % 2,
                     *csf, 700 + i if i <= 22 else None, 1.1 + i * 0.01 if i <= 33 else None,
                     ["e3/e3", "e3/e4"][i % 2] if i <= 40 else None, 0.2 if i <= 6 else None])
        if i <= 30:
            data.append([int(_pid(i)), int(_cohort(i)[0]), "V04", f"{(i % 12) + 1:02d}/2032", 21 + i % 9, 22 + i % 9, 1 + i % 2,
                         None, None, None, None, ["e3/e3", "e3/e4"][i % 2], None])
    dictionary = workbook.create_sheet("Data dictionary")
    dictionary.append(["Category", "Variable", "Description", "Code", "Decode", "Original Variable(s)",
                       "Original Dataset(s)", "Derivation Notes"])
    dictionary.append(["Family history", "fampd", "Family history of PD", 1, "Family history present", None, None, None])
    dictionary.append([None, None, None, 2, "No family history", None, None, None])  # variable name not repeated
    dictionary.append(["Cognition", "moca", "MoCA total score", None, None, "MCATOT", "MoCA", None])
    dictionary.append(["Biospecimen", "abeta", "CSF amyloid beta", None, None, "ABeta",
                       "Current Biospecimen Analysis Results", "Values from Project 159; 1700 is the upper limit"])
    dictionary.append(["Biospecimen", "ptau", "CSF phosphorylated tau", None, None, "pTau",
                       "Current Biospecimen Analysis Results", "Values from project 159"])
    dictionary.append(["Biospecimen", "IU_ABeta42_CSF", "CSF amyloid beta 42", None, None, "ABeta42",
                       "Current Biospecimen Analysis Results", "Project ID 283"])
    dictionary.append(["Biospecimen", "tau", "CSF total tau", None, None, "tTau", None, None])
    workbook.properties.created = datetime(2024, 3, 10)
    workbook.properties.modified = datetime(2024, 3, 10)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def biospecimen_summary() -> bytes:
    rows = [[125, "Synthetic CSF immunoassay panel", "Cerebrospinal Fluid", 888, 3032, 3, 10],
            [159, "Synthetic CSF reanalysis", "Cerebrospinal Fluid", 1027, 2439, 2, 6],
            [283, "Synthetic CSF and plasma panel", "Cerebrospinal Fluid", 1914, 4450, 2, 5],
            [237, "Synthetic seed amplification assay", "Cerebrospinal Fluid", 3654, 3900, 1, 4]]
    return _csv(["PROJECTID", "PROJECT_NAME", "TYPE", "UNIQUE_SUBJECT_COUNT", "UNIQUE_SUBJECT_VISIT_COUNT",
                 "AVG_VISITS_PER_SUBJECT", "MAX_VISITS_PER_SUBJECT"], rows)


def fs7_aseg() -> bytes:
    rows = [[_pid(i), "BL", 900 + i * 3, 1900 + i * 5, 12 + i, 14000 + i * 20, 4800 + i * 9, 3600 + i * 7]
            for i in range(1, 21)]
    return _csv(["PATNO", "EVENT_ID", "Left_WM_hypointensities", "WM_hypointensities", "non_WM_hypointensities",
                 "Left_Cerebellum_White_Matter", "Left_Putamen", "Left_Caudate"], rows)


def grey_matter() -> bytes:
    rows = [[_pid(i), "BL", 600000 + i * 100] for i in range(1, 11)]
    return _csv(["PATNO", "EVENT_ID", "GM_VOLUME"], rows)


def remote_demographics() -> bytes:
    return _csv(["PATNO", "SEX"], [[_pid(i), i % 2] for i in range(41, 81)])


def remote_race() -> bytes:
    return _csv(["PATNO", "EVENT_ID", "RAWHITE"], [[_pid(i), "RACE", i % 2] for i in range(55, 81)])


def adverse_events() -> bytes:
    rows = [[7000 + i, _pid(i), "LOG", 1 + i % 3] for i in range(1, 11)]
    return _csv(["REC_ID", "PATNO", "EVENT_ID", "AESEVER"], rows)


def research_diagnosis() -> bytes:
    rows = [[8000 + i, _pid(i), "BL", 1 if i <= 25 else 17, i % 2, i % 2] for i in range(1, 41)]
    return _csv(["REC_ID", "PATNO", "EVENT_ID", "PRIMDIAG", "DATSCAN", "UPDRS3"], rows)


def saa_results() -> bytes:
    rows = []
    for i in range(1, 37):
        status = ["Positive", "Negative", "Positive", "Inconclusive", "Positive", "Negative"][i % 6]
        rows.append([_pid(i), _cohort(i)[0], "BL" if i <= 30 else "V04", "Cerebrospinal Fluid", "24h", status, 237])
    return _csv(["PATNO", "COHORT", "CLINICAL_EVENT", "TYPE", "SAAMethod", "SAA_Status", "PROJECTID"], rows)


def ida_search() -> bytes:
    rows = []
    for i in range(1, 9):
        rows.append([_pid(i), "PD", "M", 60 + i, "Baseline", "SPECT", "DaTSCAN", f"I{770000 + i}"])
        rows.append([_pid(i), "PD", "M", 61 + i, "Month 12", "MRI", "T1 SYNTHTEXT sagittal", f"I{780000 + i}"])
    return _csv(["Subject ID", "Group", "Sex", "Age", "Visit", "Modality", "Description", "Image ID"], rows)


def derived_dataset() -> bytes:
    header = ["PATNO", "EVENT_ID", "NP3TOT", "MCATOT", "ABETA_42", "COHORT_x", "AD_like_label", "mystery_feature"]
    events = ["BL", "SC", "V04", "RACE", "LOG", "P102"]
    rows = []
    for i in range(1, 41):
        for k in range(3):
            rows.append([_pid(i), events[(i + k) % 6], 10 + i, "" if i % 4 == 0 else 20 + i % 9,
                         "" if i % 3 else 600 + i, _cohort(i)[0], i % 2, round(i * 0.37, 2)])
    rows += rows[:12]  # row multiplication left over from an outer join
    return _csv(header, rows)


def build(dest: Path = DEFAULT_DEST) -> Path:
    """Write the synthetic download directory to ``dest`` (replacing it) and return it."""
    dest = Path(dest)
    if dest.exists():
        shutil.rmtree(dest)
    dest.mkdir(parents=True)

    too_deep = _zip([("too_deep.csv", _csv(["PATNO", "X"], [[_pid(1), 1]]))], _STAMP_A)
    nested = _zip([
        ("iu_genetic_consensus_20240101.csv", genetics()),
        ("Legacy_Table.xls", b"synthetic placeholder, not a real workbook"),
        ("deeper.zip", too_deep),
    ], _STAMP_A)
    first = _zip([
        ("Participant_Status_01Jan2024.csv", participant_status(50)),
        ("Demographics_01Jan2024.csv", demographics()),
        ("MDS-UPDRS_Part_III_01Jan2024.csv", updrs3(extra_column=False)),
        ("Montreal_Cognitive_Assessment__MoCA__01Jan2024.csv", moca()),
        ("Current_Biospecimen_Analysis_Results_01Jan2024.csv", current_biospecimen()),
        ("Code_List_01Jan2024.csv", code_list()),
        ("Data_Dictionary_01Jan2024.csv", data_dictionary()),
        ("Data_Dictionary_-__Annotated__01Jan2024.csv", annotated_dictionary()),
        ("docs/PPMI_Data_User_Guide.pdf", b"%PDF-1.4 synthetic placeholder"),
        ("docs/notes.txt", b"synthetic placeholder notes"),
        ("extra/Genetics_bundle.zip", nested),
    ], _STAMP_A)
    second = _zip([
        ("Participant_Status_15Mar2024.csv", participant_status(60)),
        ("Demographics_01Jan2024.csv", demographics()),
        ("MDS-UPDRS_Part_III_15Mar2024.csv", updrs3(extra_column=True)),
        ("Vital_Signs_15Mar2024.csv", vital_signs()),
        ("Medical_Conditions_Log_15Mar2024.csv", medical_conditions()),
        ("Concomitant_Medication_Log_15Mar2024.csv", medications()),
        ("Pilot_Biospecimen_Analysis_Results_15Mar2024.csv", pilot_biospecimen()),
        ("DaTScan_Analysis_15Mar2024.csv", datscan()),
        ("Blood_Chemistry___Hematology_15Mar2024.csv", blood_chemistry()),
        ("PPMI_SAA_Results_15Mar2024.xlsx", saa_workbook()),
        ("PPMI_Curated_Data_Cut_Public_20240310.xlsx", curated_workbook()),
        ("Current_Biospecimen_Analysis_Results_Summary_Table_15Mar2024.csv", biospecimen_summary()),
        ("FS7_ASEG_VOL_15Mar2024.csv", fs7_aseg()),
        ("FS7_APARC_CTH_15Mar2024.csv", fs7_aparc()),
        ("Grey_Matter_Volume_15Mar2024.csv", grey_matter()),
        ("ST-Direct_Demographics_15Mar2024.csv", remote_demographics()),
        ("Race_and_Ethnicity_Question_in_myPPMI_15Mar2024.csv", remote_race()),
        ("Adverse_Event_Log_15Mar2024.csv", adverse_events()),
        ("Primary_Research_Diagnosis_15Mar2024.csv", research_diagnosis()),
        ("SAA_Biospecimen_Analysis_Results_15Mar2024.csv", saa_results()),
        ("Broken_Lines_15Mar2024.csv", b"PATNO,EVENT_ID,SCORE\n9900001,BL,1\n9900002,BL,2,unexpected,extra\n9900003,BL,3\n"),
        ("Empty_Table.csv", b""),
    ], _STAMP_B)
    imaging_entries = [(f"PPMI/{_pid(i)}/DaTSCAN/2031-01-05/I77{i}/scan_{_pid(i)}.dcm", b"\x00" * 64) for i in range(1, 7)]
    imaging_entries.append((f"PPMI/{_pid(2)}/MRI/2031-02-01/T1.nii.gz", b"\x1f\x8b" + b"\x00" * 32))
    imaging_entries.append(("idaSearch_1_15_2024.csv", ida_search()))
    imaging = _zip(imaging_entries, _STAMP_B)

    files = {
        "download.zip": (first, _STAMP_A),
        "download (1).zip": (second, _STAMP_B),
        "Imaging.zip": (imaging, _STAMP_B),
        "final_multimodal_dataset.csv": (derived_dataset(), _STAMP_B),
        "PPMI_Groups_and_Subgroups_Guidance.pdf": (b"%PDF-1.4 synthetic placeholder", _STAMP_A),
    }
    from datetime import datetime, timezone

    for name, (data, stamp) in files.items():
        path = dest / name
        path.write_bytes(data)
        seconds = datetime(*stamp, tzinfo=timezone.utc).timestamp()
        os.utime(path, (seconds, seconds))
    return dest


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_DEST
    print(build(target))
