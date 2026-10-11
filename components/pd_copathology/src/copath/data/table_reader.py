"""Chunked, text-preserving readers for tabular members.

Every cell is read as text with no missing-value conversion, so that blank
cells, NA-like words and assay tokens such as ``<LLOQ`` stay distinguishable.
"""

from __future__ import annotations

import shutil
import tempfile
from typing import IO, Iterator

import pandas as pd

SUPPORTED_EXTS = {".csv", ".tsv", ".xlsx", ".xlsm"}


class UnsupportedTableFormat(Exception):
    """Raised for a tabular file type the audit cannot read."""


def _clean_columns(columns: list[object]) -> list[str]:
    names, seen = [], {}
    for idx, col in enumerate(columns):
        name = "" if col is None else str(col).strip()
        if not name or name.lower().startswith("unnamed:"):
            name = f"unnamed_{idx}"
        if name in seen:
            seen[name] += 1
            name = f"{name}.{seen[name]}"
        else:
            seen[name] = 0
        names.append(name)
    return names


def _iter_csv(stream: IO[bytes], sep: str, chunk_rows: int, max_rows: int | None, strict: bool):
    reader = pd.read_csv(
        stream,
        sep=sep,
        dtype=str,
        keep_default_na=False,
        na_filter=False,
        chunksize=chunk_rows,
        nrows=max_rows,
        encoding="utf-8-sig",
        encoding_errors="replace",
        on_bad_lines="error" if strict else "skip",
    )
    with reader:
        for chunk in reader:
            chunk.columns = _clean_columns(list(chunk.columns))
            yield None, chunk


def _iter_xlsx(stream: IO[bytes], chunk_rows: int, max_rows: int | None):
    from openpyxl import load_workbook

    # openpyxl needs a seekable file; ZIP member streams are not reliably seekable.
    with tempfile.SpooledTemporaryFile(max_size=64 * 1024 * 1024) as tmp:
        shutil.copyfileobj(stream, tmp)
        tmp.seek(0)
        workbook = load_workbook(tmp, read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                rows = sheet.iter_rows(values_only=True)
                header = None
                for row in rows:
                    if any(cell is not None and str(cell).strip() for cell in row):
                        header = _clean_columns(list(row))
                        break
                if header is None:
                    continue
                buffer, total = [], 0
                for row in rows:
                    cells = ["" if c is None else str(c) for c in row[: len(header)]]
                    cells += [""] * (len(header) - len(cells))
                    buffer.append(cells)
                    total += 1
                    if len(buffer) >= chunk_rows:
                        yield sheet.title, pd.DataFrame(buffer, columns=header, dtype=str)
                        buffer = []
                    if max_rows is not None and total >= max_rows:
                        break
                if buffer or total == 0:
                    yield sheet.title, pd.DataFrame(buffer, columns=header, dtype=str)
        finally:
            workbook.close()


def iter_table_chunks(
    stream: IO[bytes],
    ext: str,
    chunk_rows: int = 100_000,
    max_rows: int | None = None,
    strict: bool = True,
) -> Iterator[tuple[str | None, pd.DataFrame]]:
    """Yield ``(sheet_name, chunk)`` pairs; ``sheet_name`` is ``None`` for delimited text.

    With ``strict=True`` a malformed delimited line raises
    ``pandas.errors.ParserError``; with ``strict=False`` such lines are skipped.
    """
    ext = ext.lower()
    if ext == ".csv":
        yield from _iter_csv(stream, ",", chunk_rows, max_rows, strict)
    elif ext == ".tsv":
        yield from _iter_csv(stream, "\t", chunk_rows, max_rows, strict)
    elif ext in {".xlsx", ".xlsm"}:
        yield from _iter_xlsx(stream, chunk_rows, max_rows)
    else:
        raise UnsupportedTableFormat(ext)
